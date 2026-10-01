"""Turn an uploaded file into numbered pages of text.

Page numbers are PDF pages, spreadsheet sheets, or slides. A text file,
a Word document, and an image are a single page.
"""

from io import BytesIO
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from app.services.ingest import PermanentIngestError
from app.services.language import LanguageModel, TransientLanguageError

MAX_EXTRACT_CHARS = 100_000


def empty_text_error(filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return "This PDF has no extractable text. Scanned pages need OCR."
    if suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        return "This image has no readable text."
    return "This file has no extractable text."


def extract_document(filename: str, data: bytes, model: LanguageModel) -> list[tuple[int, str]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        pages = _pdf_pages(data)
    elif suffix in {".txt", ".md", ".csv"}:
        pages = [(1, data.decode("utf-8"))]
    elif suffix == ".xlsx":
        pages = _workbook_pages(data)
    elif suffix == ".docx":
        pages = _word_pages(data)
    elif suffix == ".pptx":
        pages = _slide_pages(data)
    elif suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        pages = [(1, _transcribe(model, data, suffix))]
    else:
        raise PermanentIngestError("This file type cannot be indexed.")
    return _cap(pages)


def _cap(pages: list[tuple[int, str]]) -> list[tuple[int, str]]:
    kept: list[tuple[int, str]] = []
    used = 0
    for number, text in pages:
        if used >= MAX_EXTRACT_CHARS:
            break
        clipped = text[: MAX_EXTRACT_CHARS - used]
        kept.append((number, clipped))
        used += len(clipped)
    return kept


def _pdf_pages(data: bytes) -> list[tuple[int, str]]:
    try:
        reader = PdfReader(BytesIO(data))
    except Exception as exc:
        raise PermanentIngestError("This file is not a readable PDF.") from exc
    if reader.is_encrypted:
        raise PermanentIngestError("This PDF is encrypted and cannot be indexed.")
    pages: list[tuple[int, str]] = []
    try:
        for index, page in enumerate(reader.pages, start=1):
            pages.append((index, page.extract_text() or ""))
    except Exception as exc:
        raise PermanentIngestError("This file is not a readable PDF.") from exc
    return pages


def _workbook_pages(data: bytes) -> list[tuple[int, str]]:
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False)
    except Exception as exc:
        raise PermanentIngestError("This workbook could not be read.") from exc
    pages: list[tuple[int, str]] = []
    try:
        for index, sheet in enumerate(workbook.worksheets, start=1):
            rows: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                cells = []
                for value in row:
                    if value is None:
                        continue
                    text = str(value).strip()
                    if not text or text.startswith("="):
                        continue
                    cells.append(text)
                if cells:
                    rows.append(" | ".join(cells))
            pages.append((index, "\n".join(rows)))
    finally:
        workbook.close()
    return pages


def _word_pages(data: bytes) -> list[tuple[int, str]]:
    try:
        document = Document(BytesIO(data))
    except Exception as exc:
        raise PermanentIngestError("This document could not be read.") from exc
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append("\t".join(cells))
    return [(1, "\n\n".join(parts))]


def _slide_pages(data: bytes) -> list[tuple[int, str]]:
    try:
        presentation = Presentation(BytesIO(data))
    except Exception as exc:
        raise PermanentIngestError("This presentation could not be read.") from exc
    pages: list[tuple[int, str]] = []
    for index, slide in enumerate(presentation.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = shape.text_frame.text.strip()
            if text:
                lines.append(text)
        pages.append((index, "\n".join(lines)))
    return pages


def _transcribe(model: LanguageModel, data: bytes, suffix: str) -> str:
    mime = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp"}[suffix]
    try:
        return model.transcribe(data, mime).strip()
    except TransientLanguageError:
        raise
    except Exception as exc:
        raise PermanentIngestError("This image could not be transcribed.") from exc
