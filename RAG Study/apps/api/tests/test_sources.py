import base64
import uuid
from io import BytesIO

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from sqlalchemy import select

from app.models import DocumentChunk, DocumentVersion
from app.worker import run_once
from tests.test_auth import register
from tests.test_courses import _create_course
from tests.test_documents import _upload
from tests.test_ingest import LECTURE

PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def _xlsx() -> bytes:
    workbook = Workbook()
    workbook.active.title = "Concepts"
    workbook.active["A1"] = LECTURE
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _docx() -> bytes:
    document = Document()
    document.add_paragraph(LECTURE)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pptx() -> bytes:
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1.5))
    shape.text_frame.text = LECTURE
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _chunks(db, document_id: str) -> list[DocumentChunk]:
    return list(
        db.scalars(
            select(DocumentChunk)
            .where(DocumentChunk.document_id == uuid.UUID(document_id))
            .order_by(DocumentChunk.chunk_index)
        ).all()
    )


def test_text_workbook_document_deck_and_image_are_indexed(client, db):
    register(client)
    course_id = _create_course(client).json()["id"]
    uploads = [
        ("notes.txt", LECTURE.encode(), 1, "text/plain"),
        ("sheet.xlsx", _xlsx(), 1, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        ("notes.docx", _docx(), 1, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        ("slides.pptx", _pptx(), 2, "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
        ("board.png", PNG, 1, "image/png"),
    ]
    for filename, content, page, mime in uploads:
        uploaded = _upload(client, course_id, content=content, filename=filename, title=filename)
        assert uploaded.status_code == 201, uploaded.text
        body = uploaded.json()
        assert body["mime_type"] == mime
        assert run_once(db) is True
        listed = client.get(f"/courses/{course_id}/documents").json()
        match = next(item for item in listed if item["id"] == body["id"])
        assert match["status"] == "READY"
        chunks = _chunks(db, body["id"])
        assert chunks
        assert chunks[0].page_start == page
        assert "virtual memory" in chunks[0].text.lower()
        version = db.get(DocumentVersion, uuid.UUID(body["version_id"]))
        assert version.storage_key.endswith(filename[filename.rfind(".") :])


def test_rejected_extension(client):
    register(client)
    course_id = _create_course(client).json()["id"]
    rejected = _upload(client, course_id, content=b"PK\x03\x04macro", filename="grades.xlsm")
    assert rejected.status_code == 400
